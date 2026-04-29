"""Drafting templates and export utilities."""

from __future__ import annotations

import json
import re
from io import BytesIO
from functools import lru_cache
from pathlib import Path
from typing import Any

from docx import Document

from app.config import get_settings
from app.services.ai import translate_text


def _normalize_language(value: str | None) -> str:
    if not value:
        return "en"
    lang = value.lower()
    if lang in {"en", "hi", "bilingual"}:
        return lang
    return "en"


@lru_cache()
def _load_templates() -> dict[str, Any]:
    settings = get_settings()
    path = Path(settings.FEATURE_DATA_DIR) / "drafting_templates.json"
    if not path.exists():
        return {"templates": []}
    return json.loads(path.read_text(encoding="utf-8"))


def _find_template(template_id: str, language: str) -> dict[str, Any] | None:
    data = _load_templates()
    for template in data.get("templates", []):
        if str(template.get("id")) == template_id and template.get("language") == language:
            return template
    return None


def get_template(template_id: str, language: str) -> dict[str, Any]:
    lang = _normalize_language(language)
    base = _find_template(template_id, "en")
    if not base:
        raise ValueError("Template not found")

    placeholders = base.get("placeholders", [])
    if lang == "bilingual":
        hi_template = _find_template(template_id, "hi")
        content_hi = hi_template.get("content") if hi_template else None
        generated = False
        if not content_hi:
            content_hi = translate_text(base.get("content", ""), "en", "hi")
            generated = True
        return {
            "template_id": template_id,
            "title": base.get("title"),
            "language": lang,
            "content_en": base.get("content"),
            "content_hi": content_hi,
            "placeholders": placeholders,
            "generated": generated,
        }

    if lang == "hi":
        hi_template = _find_template(template_id, "hi")
        if hi_template:
            return {
                "template_id": template_id,
                "title": hi_template.get("title"),
                "language": lang,
                "content": hi_template.get("content"),
                "placeholders": placeholders,
                "generated": False,
            }
        translated = translate_text(base.get("content", ""), "en", "hi")
        return {
            "template_id": template_id,
            "title": base.get("title"),
            "language": lang,
            "content": translated,
            "placeholders": placeholders,
            "generated": True,
        }

    return {
        "template_id": template_id,
        "title": base.get("title"),
        "language": lang,
        "content": base.get("content"),
        "placeholders": placeholders,
        "generated": False,
    }


def export_docx(title: str, content: str) -> tuple[str, BytesIO]:
    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    for line in (content or "").splitlines():
        if line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")

    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)

    safe_name = re.sub(r"[^a-zA-Z0-9_-]+", "_", title.strip() or "draft")
    return safe_name, stream
